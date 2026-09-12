"""Generate original, shareable verification documents; no private corpus."""
import argparse
from pathlib import Path
import fitz
from docx import Document
from docx.shared import Inches
from docx.oxml import parse_xml

def make_pdf(path, marker):
    document = fitz.open()
    page = document.new_page(width=595, height=842)
    page.insert_text((50, 65), marker, fontsize=23)
    page.insert_text((50, 108), 'A clear source supports a reliable knowledge base.', fontsize=14)
    page.insert_text((50, 140), 'Distance = 6 m. Time = 3 s. Speed = 2 m/s.', fontsize=14)
    for x in (50, 230, 410):
        page.draw_line((x, 190), (x, 310), width=1)
    for y in (190, 230, 270, 310):
        page.draw_line((50, y), (410, y), width=1)
    for x, y, text in [(65,215,'Quantity'),(245,215,'Value'),(65,255,'Distance'),(245,255,'6 m'),(65,295,'Time'),(245,295,'3 s')]:
        page.insert_text((x, y), text, fontsize=13)
    page.draw_line((110, 500), (390, 500), width=2)
    page.draw_line((110, 500), (110, 360), width=2)
    page.draw_line((110, 500), (340, 390), width=2)
    page.insert_text((397, 505), 't', fontsize=18)
    page.insert_text((95, 350), 's', fontsize=18)
    page.insert_text((200, 550), 'Figure 1. Motion graph', fontsize=13)
    document.save(path)
    pixels = page.get_pixmap(dpi=144, alpha=False)
    png = path.with_suffix('.png')
    pixels.save(png)
    document.close()
    return png

def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('output', type=Path)
    root = parser.parse_args().output.resolve()
    root.mkdir(parents=True, exist_ok=True)
    make_pdf(root / 'native.pdf', 'NATIVE PDF PUBLIC TEST')
    scan_png = make_pdf(root / 'scan-original.pdf', 'SCANNED PDF PUBLIC TEST')
    scan = fitz.open()
    page = scan.new_page(width=595, height=842)
    page.insert_image(page.rect, filename=str(scan_png))
    scan.save(root / 'scanned.pdf')
    scan.close()
    word_png = make_pdf(root / 'word-original.pdf', 'SCREENSHOT WORD PUBLIC TEST')
    doc = Document()
    doc.add_picture(str(word_png), width=Inches(6))
    doc.save(root / 'screenshot.docx')
    doc = Document()
    doc.add_heading('Native Word public test', level=1)
    doc.add_paragraph('Distance = 6 m. Time = 3 s. Speed = 2 m/s.')
    paragraph = doc.add_paragraph('A native fraction: ')
    paragraph._p.append(parse_xml('<m:oMath xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math"><m:f><m:num><m:r><m:t>x</m:t></m:r></m:num><m:den><m:r><m:t>2</m:t></m:r></m:den></m:f></m:oMath>'))
    table = doc.add_table(rows=2, cols=2)
    table.cell(0,0).text, table.cell(0,1).text = 'Quantity', 'Value'
    table.cell(1,0).text, table.cell(1,1).text = 'Distance', '6 m'
    doc.save(root / 'native.docx')
    print(root)

if __name__ == '__main__':
    main()
